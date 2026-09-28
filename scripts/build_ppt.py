"""
Generates a 5-slide Eklavya pitch deck from the measured content already
documented in docs/EKLAVYA_FULL_ARCHITECTURE.md, PROGRESS.md and README.md.

Run:
    python scripts/build_ppt.py
Output:
    Eklavya_5_Slide_Deck.pptx (repo root)
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------- Palette ----------
NAVY = RGBColor(0x0B, 0x1E, 0x33)
DARK = RGBColor(0x11, 0x2A, 0x46)
ACCENT = RGBColor(0x2E, 0xC4, 0xB6)   # teal accent
ACCENT2 = RGBColor(0xF2, 0xA6, 0x1E)  # amber accent
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xD7, 0xE1, 0xEC)
GREY = RGBColor(0x8A, 0x9B, 0xB0)
GOOD = RGBColor(0x4C, 0xD9, 0x7B)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color=NAVY):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def no_line(shape):
    shape.line.fill.background()


def textbox(slide, left, top, width, height, text, size=18, color=WHITE,
            bold=False, italic=False, align=PP_ALIGN.LEFT, font="Calibri",
            anchor=MSO_ANCHOR.TOP, line_spacing=1.0):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = font
    return tb


def bullets(slide, left, top, width, height, items, size=15, color=LIGHT,
            bold_first=False, font="Calibri", space_after=8, marker="›"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        p.line_spacing = 1.08
        if isinstance(item, tuple):
            head, rest = item
            r1 = p.add_run()
            r1.text = f"{marker}  {head}"
            r1.font.bold = True
            r1.font.size = Pt(size)
            r1.font.color.rgb = ACCENT
            r1.font.name = font
            if rest:
                r2 = p.add_run()
                r2.text = "  " + rest
                r2.font.size = Pt(size)
                r2.font.color.rgb = color
                r2.font.name = font
        else:
            r = p.add_run()
            r.text = f"{marker}  {item}"
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.name = font
    return tb


def rect(slide, left, top, width, height, fill=DARK, line=False, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_type, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if not line:
        no_line(sh)
    else:
        sh.line.color.rgb = ACCENT
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def header(slide, kicker, title, slide_no):
    rect(slide, 0, 0, SLIDE_W, Inches(1.15), fill=DARK)
    rect(slide, 0, Inches(1.15), SLIDE_W, Pt(3), fill=ACCENT)
    textbox(slide, Inches(0.5), Inches(0.12), Inches(9), Inches(0.35),
            kicker.upper(), size=13, color=ACCENT2, bold=True)
    textbox(slide, Inches(0.5), Inches(0.42), Inches(11.5), Inches(0.7),
            title, size=30, color=WHITE, bold=True)
    textbox(slide, Inches(12.4), Inches(0.1), Inches(0.7), Inches(0.4),
            f"{slide_no}/5", size=14, color=GREY, align=PP_ALIGN.RIGHT)
    textbox(slide, Inches(0.5), Inches(7.05), Inches(9), Inches(0.35),
            "Eklavya Strategy for Electronic Warfare", size=10, color=GREY)


def stat_card(slide, left, top, width, height, number, label, color=ACCENT):
    rect(slide, left, top, width, height, fill=DARK, radius=True)
    textbox(slide, left, top + Inches(0.12), width, Inches(0.6), number,
            size=26, color=color, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, left + Inches(0.1), top + height - Inches(0.55), width - Inches(0.2),
            Inches(0.5), label, size=11, color=LIGHT, align=PP_ALIGN.CENTER)


def add_table(slide, left, top, width, height, data, col_widths=None,
              header_fill=ACCENT, body_fill=DARK, font_size=12, header_size=12):
    rows = len(data)
    cols = len(data[0])
    gtable = slide.shapes.add_table(rows, cols, left, top, width, height).table
    if col_widths:
        for i, w in enumerate(col_widths):
            gtable.columns[i].width = w
    for r in range(rows):
        for c in range(cols):
            cell = gtable.cell(r, c)
            cell.text = str(data[r][c])
            cell.margin_left = Pt(6)
            cell.margin_right = Pt(6)
            cell.margin_top = Pt(3)
            cell.margin_bottom = Pt(3)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
            run = para.runs[0]
            run.font.size = Pt(header_size if r == 0 else font_size)
            run.font.bold = (r == 0)
            run.font.name = "Calibri"
            run.font.color.rgb = NAVY if r == 0 else LIGHT
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.fore_color.rgb = DARK if r % 2 == 1 else RGBColor(0x16, 0x33, 0x54)
    return gtable


# =====================================================================
# SLIDE 1 — Novelty & Key Features
# =====================================================================
s = add_slide()
set_bg(s)
header(s, "Section 1", "Novelty & Key Features", 1)

textbox(s, Inches(0.5), Inches(1.35), Inches(7.6), Inches(0.6),
        "Eklavya learns WHEN and WHERE to look in the spectrum — with no prior "
        "threat library — from its own hits and misses.", size=15, color=LIGHT, italic=True)

left_items = [
    ("Interval-censored timing inference",
     "A miss isn't a null event — a d-ms look that saw nothing eliminates a whole phase interval for every candidate period."),
    ("Learned discrete period prior",
     "96.2% of 128,252 real radar scan periods fall on just 22 values → after one catch, ~19 candidate return times, not a continuum."),
    ("Two competing emitter models",
     "Periodic flash vs. Gilbert–Elliott on/off, arbitrated by Bayes factors — covers the 88% of real emitters that aren't clean periodic flashes."),
    ("Value-per-unit-receiver-time index",
     "Unifies exploit / discover-new / resolve-timing on one scale, so a 5 ms and a 100 ms option are directly comparable."),
]
bullets(s, Inches(0.5), Inches(2.05), Inches(7.4), Inches(4.5), left_items, size=14.5)

right_items = [
    ("Micro dwell controller",
     "5 ms search looks, extend-on-partial-pulse, stop-on-confirm → lock-out 60.2%→69.1%, first-catch time cut 30–35%."),
    ("Anti-synchronization by design",
     "Golden-ratio dither in search mode — structurally immune to the lock-out that defeats fixed sweeps."),
    ("Explainable by construction",
     "Every dwell decision logs the numbers that produced it — shown live in the demo's decision panel."),
    ("Threat-aware priority (optional)",
     "Classifies emitter threat level from observed parameters alone; operator-selectable mission profile."),
]
bullets(s, Inches(8.1), Inches(2.05), Inches(4.7), Inches(4.5), right_items, size=14.5)

rect(s, Inches(0.5), Inches(6.35), Inches(12.3), Inches(0.55), fill=DARK, radius=True)
textbox(s, Inches(0.7), Inches(6.42), Inches(12), Inches(0.4),
        "Design principle: don't ask ML to rediscover radar physics — model what is known, "
        "learn the parameters online, let learning handle the residue.", size=13, color=ACCENT2, bold=True)

# =====================================================================
# SLIDE 2 — Technical Architecture
# =====================================================================
s = add_slide()
set_bg(s)
header(s, "Section 2", "Technical Architecture", 2)

textbox(s, Inches(0.5), Inches(1.35), Inches(12.3), Inches(0.4),
        "Closed loop: observe → associate → predict → score → schedule → learn", size=15, color=ACCENT, bold=True)

layers = [
    ("L0", "RF Front-End", "Antenna → tunable receiver → digitiser"),
    ("L1", "Signal Processing", "Pulse detection (CFAR) → PDW extraction"),
    ("L2", "Deinterleaving", "Gate on AoA/freq/PW → cluster into tracks"),
    ("L3", "Emitter Knowledge", "Periodic + on/off model bank, undiscovered belief, learned priors"),
    ("L4", "Decision", "Rate index, micro dwell controller, dither, staleness governor"),
    ("L5", "Learning", "Online: timing posteriors, visibility κ | Offline: priors, tuning, residual RL"),
    ("L6", "Control & HMI", "Operator priorities, decision explanations, C2/fusion interface"),
]
top = Inches(1.95)
box_h = Inches(0.62)
gap = Inches(0.08)
box_w = Inches(7.7)
for i, (code, name, desc) in enumerate(layers):
    y = top + i * (box_h + gap)
    rect(s, Inches(0.5), y, Inches(0.75), box_h, fill=ACCENT, radius=True)
    textbox(s, Inches(0.5), y + Inches(0.06), Inches(0.75), Inches(0.5), code,
            size=16, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    rect(s, Inches(1.35), y, box_w, box_h, fill=DARK, radius=True)
    textbox(s, Inches(1.5), y + Inches(0.05), Inches(2.0), Inches(0.5), name,
            size=13, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, Inches(3.55), y + Inches(0.03), box_w - Inches(2.2), box_h - Inches(0.05), desc,
            size=10.5, color=LIGHT, anchor=MSO_ANCHOR.MIDDLE)
    if i < len(layers) - 1:
        arrow = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(0.78), y + box_h - Inches(0.02),
                                    Inches(0.15), gap + Inches(0.05))
        arrow.fill.solid(); arrow.fill.fore_color.rgb = ACCENT; no_line(arrow)

# Right column: decision rule + stack
rect(s, Inches(9.35), Inches(1.95), Inches(3.5), Inches(2.35), fill=DARK, radius=True)
textbox(s, Inches(9.5), Inches(2.05), Inches(3.2), Inches(0.3), "DECISION RULE", size=12, color=ACCENT2, bold=True)
textbox(s, Inches(9.5), Inches(2.35), Inches(3.2), Inches(1.9),
        "Value per unit receiver time:\n\n"
        "I(c,d) =\n"
        "  [ exploit known emitters\n"
        "  + discover new emitters\n"
        "  + λ · information gain ]\n"
        "  ÷ (dwell + retune cost)",
        size=11.5, color=LIGHT, font="Consolas")

rect(s, Inches(9.35), Inches(4.45), Inches(3.5), Inches(2.15), fill=DARK, radius=True)
textbox(s, Inches(9.5), Inches(4.55), Inches(3.2), Inches(0.3), "TECH STACK", size=12, color=ACCENT2, bold=True)
stack_items = ["Python 3.12 + NumPy — core index/predict",
               "Gymnasium + SB3 Double-DQN — ML comparator",
               "FastAPI backend  ·  Next.js/React frontend",
               "pytest (20 tests) · seeded, replayable benchmark"]
bullets(s, Inches(9.5), Inches(4.85), Inches(3.25), Inches(1.7), stack_items, size=10.5, marker="•", space_after=6)

textbox(s, Inches(0.5), Inches(6.75), Inches(12.3), Inches(0.35),
        "Repo modules: primes/env (receiver+scenarios) · primes/predict (track models) · "
        "primes/policies/primes_lite.py (index) · api/main.py + web/ (demo)",
        size=10.5, color=GREY, italic=True)

# =====================================================================
# SLIDE 3 — Impact & Benefits
# =====================================================================
s = add_slide()
set_bg(s)
header(s, "Section 3", "Impact & Benefits", 3)

stat_card(s, Inches(0.5), Inches(1.5), Inches(2.95), Inches(1.3), "4.2×", "more flashes intercepted vs.\nopen-loop sweep (real data)", color=ACCENT)
stat_card(s, Inches(3.6), Inches(1.5), Inches(2.95), Inches(1.3), "69.8%", "Eklavya interception on real\nTuring dataset scenarios", color=GOOD)
stat_card(s, Inches(6.7), Inches(1.5), Inches(2.95), Inches(1.3), "95%", "lock-out emitters discovered\n(raster: only 15%)", color=ACCENT2)
stat_card(s, Inches(9.8), Inches(1.5), Inches(2.95), Inches(1.3), "3.2×", "faster time-to-first-catch\n(lock-out: 55.7s → 17.2s)", color=ACCENT)

textbox(s, Inches(0.5), Inches(3.0), Inches(6.5), Inches(0.35), "HEADLINE BENCHMARK (610 runs, real dataset scenarios)",
        size=13, color=ACCENT2, bold=True)
data = [
    ["Scheme", "Interception"],
    ["Open-loop raster", "16.8%"],
    ["Greedy bandit", "20.3%"],
    ["Double-DQN", "15.3%"],
    ["Eklavya", "69.8%"],
    ["Clairvoyant oracle (upper bound)", "84.8%"],
]
add_table(s, Inches(0.5), Inches(3.4), Inches(6.5), Inches(2.9), data,
          col_widths=[Inches(4.3), Inches(2.2)])

op_items = [
    ("Operational", "4.2× more of the EM picture; blind spots eliminated; earlier warning; better PRI/scan estimates from more samples."),
    ("Economic", "Zero new hardware — a software upgrade to fielded receivers; one artefact set deploys fleet-wide."),
    ("Strategic", "Works without an imported threat library; priors refit locally; transparent, auditable AI (not black-box RL)."),
    ("Beyond defence", "Spectrum regulation, interference hunting, radio-astronomy RFI mitigation, general sensor scheduling."),
]
bullets(s, Inches(7.35), Inches(3.4), Inches(5.5), Inches(3.4), op_items, size=13)

# =====================================================================
# SLIDE 4 — Feasibility & Viability
# =====================================================================
s = add_slide()
set_bg(s)
header(s, "Section 4", "Feasibility & Viability", 4)

textbox(s, Inches(0.5), Inches(1.35), Inches(12.3), Inches(0.4),
        "Not a proposal — a working, tested prototype with measured results.", size=15, color=ACCENT, bold=True)

evidence = [
    "20 unit tests · 610-run multi-seed benchmark (95% bootstrap CIs)",
    "1,008-run robustness sweep across P_d / P_fa conditions",
    "Validated on 11 real Turing dataset scenarios, not just simulation",
    "Eklavya / oracle ratio stays 0.67–0.75 as detection probability drops 1.0→0.5",
    "≤2-point interception loss up to ~23 false alarms/run",
]
bullets(s, Inches(0.5), Inches(1.95), Inches(6.1), Inches(2.4), evidence, size=13.5)

rect(s, Inches(0.5), Inches(4.5), Inches(6.1), Inches(2.15), fill=DARK, radius=True)
textbox(s, Inches(0.65), Inches(4.6), Inches(5.8), Inches(0.3), "RESOURCE FEASIBILITY", size=12, color=ACCENT2, bold=True)
res_items = [
    "Dev compute: a laptop (benchmark ~45 min on 3 cores)",
    "Data: open, Apache-2.0 licensed, streamed (no bulk storage)",
    "Licences/cost: $0 — fully open-source stack",
    "Runtime target: ~10⁴ ops/decision → fits embedded CPU",
]
bullets(s, Inches(0.65), Inches(4.9), Inches(5.75), Inches(1.7), res_items, size=12, marker="•", space_after=5)

# TRL roadmap table
textbox(s, Inches(6.95), Inches(1.95), Inches(5.9), Inches(0.35), "PATH TO FIELDED CAPABILITY", size=12, color=ACCENT2, bold=True)
trl = [
    ["Stage", "Work", "Effort"],
    ["TRL 3–4 (now)", "Algorithm + simulator + dataset validation + demo", "Done"],
    ["TRL 4–5", "Real deinterleaving, HW-in-loop (SDR), C++/ONNX port", "3–6 mo"],
    ["TRL 5–6", "ESM suite integration, multi-receiver, trials", "6–12 mo"],
    ["TRL 6+", "Multi-platform fusion, certification, field trials", "12+ mo"],
]
add_table(s, Inches(6.95), Inches(2.35), Inches(5.9), Inches(2.4), trl,
          col_widths=[Inches(1.5), Inches(3.2), Inches(1.2)], font_size=11, header_size=11)

rect(s, Inches(6.95), Inches(4.95), Inches(5.9), Inches(1.7), fill=DARK, radius=True)
textbox(s, Inches(7.1), Inches(5.05), Inches(5.6), Inches(0.3), "KEY RISK & MITIGATION", size=12, color=ACCENT2, bold=True)
risk_items = [
    ("Missed detections/false alarms", "Graceful degradation is measured, not assumed."),
    ("Real deinterleaving imperfect", "Prototype assumes upstream ID; dataset benchmarks this gap directly."),
    ("Embedded real-time budget", "Bounded per-decision cost by design; ONNX/C++ port path defined."),
]
bullets(s, Inches(7.1), Inches(5.35), Inches(5.6), Inches(1.25), risk_items, size=11, space_after=4)

# =====================================================================
# SLIDE 5 — Links & References
# =====================================================================
s = add_slide()
set_bg(s)
header(s, "Section 5", "Links & References", 5)

rect(s, Inches(0.5), Inches(1.55), Inches(6.1), Inches(1.85), fill=DARK, radius=True)
textbox(s, Inches(0.65), Inches(1.65), Inches(5.8), Inches(0.3), "PROJECT LINKS", size=13, color=ACCENT2, bold=True)
link_items = [
    "Repo: github.com/abhishek-pandey7/Smart-Scan-strategy-for-Electronic-Warfare",
    "Full architecture doc: docs/EKLAVYA_FULL_ARCHITECTURE.md",
    "Live demo (local): localhost:3000/mission?scenario=lockout",
]
bullets(s, Inches(0.65), Inches(2.0), Inches(5.8), Inches(1.4), link_items, size=12, marker="•", space_after=6)

rect(s, Inches(0.5), Inches(3.55), Inches(6.1), Inches(3.05), fill=DARK, radius=True)
textbox(s, Inches(0.65), Inches(3.65), Inches(5.8), Inches(0.3), "REPRODUCE IT YOURSELF", size=13, color=ACCENT2, bold=True)
repro = (
    "pip install -r requirements.txt\n"
    "pytest -q                                  # 20 tests\n"
    "python -m experiments.benchmark --seeds 10 # rebuild benchmark.json\n"
    "uvicorn api.main:app --port 8000            # backend\n"
    "cd web && npm install && npm run dev        # demo UI"
)
textbox(s, Inches(0.65), Inches(4.0), Inches(5.8), Inches(2.4), repro, size=12, color=GOOD, font="Consolas", line_spacing=1.3)

textbox(s, Inches(6.95), Inches(1.55), Inches(5.9), Inches(0.35), "PRIOR ART WE BUILD ON", size=13, color=ACCENT2, bold=True)
refs = [
    "Clarkson, Optimisation of periodic search strategies for ES, IEEE Trans. AES 2011 — lock-out theory",
    "Liu & Zhao, Indexability of restless bandits & Whittle index, IEEE Trans. IT 2010 — index paradigm",
    "Nakhleh et al., NeurWIN: Neural Whittle Index Network, NeurIPS 2021 — learned index roadmap",
    "ML-based Receiver Scheduling for ES (Double-DQN), IEEE Xplore 11310649 — our comparator baseline",
    "The Alan Turing Institute, Turing Synthetic Radar Dataset, arXiv:2602.03856 — validation data",
]
bullets(s, Inches(6.95), Inches(1.95), Inches(5.9), Inches(3.2), refs, size=11, marker="•", space_after=8)

rect(s, Inches(6.95), Inches(5.35), Inches(5.9), Inches(1.25), fill=DARK, radius=True)
textbox(s, Inches(7.1), Inches(5.45), Inches(5.6), Inches(1.0),
        "Every figure in this deck is reproducible bit-for-bit from a seed — "
        "no number here is hand-picked.", size=12, color=ACCENT2, italic=True, bold=True)

prs.save("Eklavya_5_Slide_Deck.pptx")
print("Saved Eklavya_5_Slide_Deck.pptx")
